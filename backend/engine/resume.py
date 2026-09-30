"""Turn a resume (txt / docx / pdf) into numbered evidence items.

An evidence item is one defensible statement from the resume:
    {"id": "E3", "section": "Experience", "context": "Data Analyst Intern, Northwind",
     "text": "Built a Tableau dashboard ...", "skills": [...], "numbers": [...]}
"""
from __future__ import annotations

import io
import re

from .skills import extract_skills
from .text import extract_numbers, is_bullet, normalize_ws, split_lines, strip_bullet

SECTION_ALIASES = {
    "Summary": ["summary", "profile", "objective", "professional summary", "about me"],
    "Education": ["education", "academic background", "academics"],
    "Experience": ["experience", "work experience", "professional experience", "employment",
                   "internships", "internship", "work history", "relevant experience"],
    "Projects": ["projects", "academic projects", "personal projects", "selected projects", "project experience"],
    "Research": ["research", "research experience", "publications", "papers"],
    "Skills": ["skills", "technical skills", "core skills", "technologies", "tools", "skills & tools", "technical expertise"],
    "Certifications": ["certifications", "certificates", "licenses", "licenses & certifications", "certifications & training"],
    "Leadership": ["leadership", "activities", "extracurricular", "volunteering", "volunteer experience", "involvement"],
    "Awards": ["awards", "honors", "achievements", "honors & awards"],
}
# Skills only count as "shown in work" when they appear in these sections.
DEMONSTRATED = {"Experience", "Projects", "Research", "Leadership", "Awards"}
_HEADING_LOOKUP = {a: canon for canon, aliases in SECTION_ALIASES.items() for a in aliases}

CONTACT_RE = re.compile(r"@|linkedin|github\.com|https?://|\+?\d[\d\s().-]{8,}\d")
DATE_RE = re.compile(r"(?i)\b\d{1,2}/\d{4}\b|\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)[a-z]*\.?\s*\d{4}|\b(19|20)\d{2}\b|present|current|expected")
ACTION_START = re.compile(r"(?i)^(built|developed|designed|created|implemented|led|managed|analy[sz]ed|cleaned|trained|"
                          r"deployed|wrote|automated|improved|reduced|increased|tutored|taught|held|documented|"
                          r"conducted|collaborated|compared|engineered|optimi[sz]ed|researched|presented|"
                          r"supported|assisted|maintained|migrated|integrated|configured|coordinated|organized|"
                          r"prepared|evaluated|tested|launched|scraped|visuali[sz]ed|modeled|annotated|explained)\b")


# --------------------------------------------------------------------------- file reading
def read_file(filename: str, data: bytes) -> str:
    name = (filename or "").lower()
    if name.endswith(".docx"):
        import docx  # python-docx
        d = docx.Document(io.BytesIO(data))
        lines = []
        for p in d.paragraphs:
            t = p.text.strip()
            style = (p.style.name or "").lower() if p.style is not None else ""
            if t and ("list" in style or p._p.pPr is not None and p._p.pPr.numPr is not None):
                t = "- " + t
            lines.append(t)
        for table in d.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells if c.text.strip()))
        return "\n".join(lines)
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((pg.extract_text() or "") for pg in reader.pages)
    if name.endswith((".txt", ".md")):
        return data.decode("utf-8", errors="replace")
    raise ValueError("Unsupported file type. Upload a .docx, .pdf or .txt resume.")


# --------------------------------------------------------------------------- parsing
HEADING_KEYWORDS = [  # order matters: first match wins for combined headings
    ("Experience", r"experience|employment|work history|internship"),
    ("Projects", r"project"),
    ("Research", r"research|publication"),
    ("Skills", r"skill|technolog|tools|competenc"),
    ("Certifications", r"certif|licen"),
    ("Education", r"education|academic"),
    ("Leadership", r"activit|leadership|volunteer|involvement|extracurricular"),
    ("Awards", r"award|honou?r|accomplish|achievement"),
    ("Summary", r"summary|profile|objective|about"),
]


def _heading(line: str) -> str | None:
    raw = normalize_ws(line).strip(":").strip()
    if not raw or len(raw) > 48 or is_bullet(line):
        return None
    key = re.sub(r"[^a-z& ]", "", raw.lower()).strip()
    if key in _HEADING_LOOKUP:
        return _HEADING_LOOKUP[key]
    # ALL-CAPS short line with a known keyword, e.g. "TECHNICAL SKILLS AND CERTIFICATIONS"
    letters = re.sub(r"[^A-Za-z]", "", raw)
    if letters and letters.isupper() and len(raw.split()) <= 6 and not re.search(r"\d", raw):
        for canon, pat in HEADING_KEYWORDS:
            if re.search(pat, key):
                return canon
    return None


def _looks_like_context(line: str, section: str) -> bool:
    """Role / project title lines, e.g. 'Data Analyst Intern, Northwind (Jan 2025 – Jun 2025)'."""
    if is_bullet(line) or ACTION_START.match(line):
        return False
    if section in ("Experience", "Research", "Leadership", "Awards") and (DATE_RE.search(line) or len(line) < 90):
        return len(line.split()) <= 16
    if section == "Projects":
        return len(line.split()) <= 14 and not line.endswith(".")
    return False


def parse_resume(text: str) -> dict:
    lines = [normalize_ws(l) for l in split_lines(text)]
    name = next((l for l in lines if l and not CONTACT_RE.search(l)), "")
    section, context = "Summary", ""
    items: list[dict] = []
    started = False
    pending: dict | None = None
    stack = ""
    last_was_context = False

    def flush():
        nonlocal pending
        if pending and len(pending["text"]) >= 3:
            items.append(pending)
        pending = None

    raw_lines = split_lines(text)
    for raw, line in zip(raw_lines, lines):
        if not line:
            flush()
            continue
        h = _heading(line)
        if h:
            flush()
            section, context, started = h, "", True
            continue
        if not started:
            continue  # name / contact block before the first heading
        if CONTACT_RE.search(line) and not ACTION_START.match(strip_bullet(line)):
            continue
        bullet = is_bullet(raw)
        body = strip_bullet(raw).strip() if bullet else line
        if last_was_context and not bullet and len(extract_skills(line)) >= 2 and ("," in line or "|" in line) \
                and not ACTION_START.match(line):
            stack = line  # "Python, Flask, spaCy | Synthetic Data" under a project title
            last_was_context = False
            continue
        if section in ("Experience", "Projects", "Research", "Leadership") and not bullet and _looks_like_context(line, section):
            flush()
            context, stack, last_was_context = line, "", True
            continue
        last_was_context = False
        # continuation of a wrapped bullet line
        if pending and not bullet and pending["section"] == section and not pending["text"].endswith(".") \
                and body[:1].islower():
            pending["text"] += " " + body
            continue
        flush()
        pending = {"section": section, "context": context, "text": body, "stack": stack}
    flush()

    evidence = []
    for i, it in enumerate(items, 1):
        full = f"{it['context']} {it['text']}"
        evidence.append({
            "id": f"E{i}",
            "section": it["section"],
            "context": it["context"],
            "text": it["text"],
            "skills": extract_skills(it["text"] if it["section"] != "Skills" else it["text"]),
            "context_skills": extract_skills(f"{it['context']} {it.get('stack', '')}") if it["context"] else [],
            "stack": it.get("stack", ""),
            "numbers": extract_numbers(it["text"]),
            "demonstrated": it["section"] in DEMONSTRATED,
        })
        _ = full
    return {"name": name if len(name.split()) <= 5 else "", "evidence": evidence}


def enrich(item: dict) -> dict:
    """Recompute derived fields after a user edits an evidence item."""
    item["text"] = normalize_ws(item.get("text", ""))
    item["context"] = normalize_ws(item.get("context", ""))
    item["skills"] = extract_skills(item["text"])
    item["context_skills"] = extract_skills(item["context"]) if item["context"] else []
    item["numbers"] = extract_numbers(item["text"])
    item["demonstrated"] = item.get("section") in DEMONSTRATED
    return item


def renumber(evidence: list[dict]) -> list[dict]:
    for i, e in enumerate(evidence, 1):
        e["id"] = f"E{i}"
    return evidence


def all_skills(evidence: list[dict]) -> dict[str, list[str]]:
    """skill -> evidence ids that mention it."""
    out: dict[str, list[str]] = {}
    for e in evidence:
        for s in e["skills"] + e.get("context_skills", []):
            out.setdefault(s, [])
            if e["id"] not in out[s]:
                out[s].append(e["id"])
    return out
